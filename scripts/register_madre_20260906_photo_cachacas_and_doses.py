#!/usr/bin/env python3
"""Cadastra cachacas das fotos 2026-09-06 e cria fichas de dose 60 ml na Madre."""

from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


API_BASE = "https://gestor-estoque-zqw9.onrender.com/api"
COMPANY_ID = 13
PRICE_WORKBOOK = Path("/home/leomassoni/Downloads/Lista de preços Bar.xlsx")
OUTPUT_WORKBOOK = Path("/home/leomassoni/Downloads/Bebidas novas fotos 2026-09-06 - custos e doses Casa de mi Madre.xlsx")
BACKUP_ROOT = Path("backups")
AUDIT_ROOT = Path("auditorias")
DOSE_VOLUME_ML = Decimal("60")
MAX_CMV = Decimal("35")
EUR_BRL = Decimal("5.95496")
IOF_PERCENT = Decimal("3.5")


PHOTO_ITEMS: list[dict[str, Any]] = [
    {
        "photo": "WhatsApp Image 2026-09-06 at 11.42.21(3).jpeg",
        "product_name": "CACHACA TIE PRATA",
        "volume_ml": Decimal("700"),
        "alcohol": "42",
        "family": "BEBIDAS",
        "subfamily": "CACHACAS",
        "currency": "BRL",
        "source_value": Decimal("76.35"),
        "exchange_rate": Decimal("1"),
        "iof_percent": Decimal("0"),
        "cost_source": "Media online: Cachacaria Nacional R$81,35; Emporio in Casa R$71,90; Emporio Bem Minerim R$74,00; Magazine Luiza/Bondfaro R$78,13.",
        "data_source": "Rotulo da foto: Tiê Prata, 700 ml, 42% vol.",
        "aliases": ["CACHACA TIE PRATA", "TIE PRATA"],
    },
    {
        "photo": "WhatsApp Image 2026-09-06 at 11.42.21(2).jpeg",
        "product_name": "CACHACA TIE CASTANHEIRA",
        "volume_ml": Decimal("700"),
        "alcohol": "40",
        "family": "BEBIDAS",
        "subfamily": "CACHACAS",
        "currency": "EUR",
        "source_value": Decimal("34.90"),
        "exchange_rate": EUR_BRL,
        "iof_percent": IOF_PERCENT,
        "cost_source": "Frau Cachaca: EUR 34,90; sem preco nacional localizado; convertido por EUR/BRL 5,95496 + 3,5% IOF.",
        "data_source": "Rotulo da foto: Tiê Castanheira, 700 ml, 40% vol.; fonte web confirma produto.",
        "aliases": ["CACHACA TIE CASTANHEIRA", "TIE CASTANHEIRA"],
    },
    {
        "photo": "WhatsApp Image 2026-09-06 at 11.42.21(1).jpeg",
        "product_name": "CACHACA MEIA LUA SALINAS BALSAMO",
        "volume_ml": Decimal("670"),
        "alcohol": "41",
        "family": "BEBIDAS",
        "subfamily": "CACHACAS",
        "currency": "BRL",
        "source_value": Decimal("47.72"),
        "exchange_rate": Decimal("1"),
        "iof_percent": Decimal("0"),
        "cost_source": "Media online: Cachacas de Salinas R$46,25; Galeria da Cachaca R$49,00; Pe de Cuba R$47,90.",
        "data_source": "Rotulo da foto e Cachacas de Salinas: Meia Lua 670 ml, balsamo, 41% vol.",
        "aliases": ["CACHACA MEIA LUA", "MEIA LUA SALINAS", "MEIA LUA BALSAMO"],
        "barcode": "7898054370162",
    },
    {
        "photo": "WhatsApp Image 2026-09-06 at 11.42.21.jpeg",
        "product_name": "CACHACA COM JAMBU SENSACAO TRADICIONAL",
        "volume_ml": Decimal("750"),
        "alcohol": "34",
        "family": "BEBIDAS",
        "subfamily": "CACHACAS",
        "currency": "BRL",
        "source_value": Decimal("69.29"),
        "exchange_rate": Decimal("1"),
        "iof_percent": Decimal("0"),
        "cost_source": "Media online sem outliers: Emporio da Raposa R$67,41; Banca do Ramon R$74,90/R$71,16 PIX.",
        "data_source": "Rotulo da foto e Emporio da Raposa: Cachaça com Jambu Sensação Tradicional, 750 ml.",
        "aliases": ["CACHACA COM JAMBU SENSACAO", "JAMBU SENSACAO"],
    },
    {
        "photo": "WhatsApp Image 2026-09-06 at 11.42.20(1).jpeg",
        "product_name": "CACHACA CAPUEIRA OURO BALSAMO",
        "volume_ml": Decimal("700"),
        "alcohol": "38",
        "family": "BEBIDAS",
        "subfamily": "CACHACAS",
        "currency": "BRL",
        "source_value": Decimal("65.00"),
        "exchange_rate": Decimal("1"),
        "iof_percent": Decimal("0"),
        "cost_source": "Valor informado pelo usuario em 2026-09-05; pesquisas online vistas entre R$79,90 e R$80,00.",
        "data_source": "Rotulo da foto: Cachaça Capueira Ouro Balsamo, 700 ml; teor lido como 38% vol.",
        "aliases": ["CACHACA CAPUEIRA OURO", "CAPUEIRA OURO BALSAMO"],
    },
    {
        "photo": "WhatsApp Image 2026-09-06 at 11.42.20.jpeg",
        "product_name": "FAMIGERADA JENI LIQUOR LICOR DE JENIPAPO",
        "volume_ml": Decimal("500"),
        "alcohol": "25",
        "family": "BEBIDAS",
        "subfamily": "LICORES",
        "currency": "BRL",
        "source_value": Decimal("99.00"),
        "exchange_rate": Decimal("1"),
        "iof_percent": Decimal("0"),
        "cost_source": "Valor informado pelo usuario em 2026-09-05; pesquisas online vistas entre R$93,02 e R$139,90.",
        "data_source": "Rotulo da foto e Amazon: Famigerada Jeni Liquor Licor de Jenipapo, 500 ml, 25% vol.",
        "aliases": ["FAMIGERADA JENI", "JENI LIQUOR", "LICOR DE JENIPAPO FAMIGERADA"],
    },
    {
        "photo": "WhatsApp Image 2026-09-06 at 11.42.19(2).jpeg",
        "product_name": "CACHACA MATRIARCA JAQUEIRA",
        "volume_ml": Decimal("750"),
        "alcohol": "42",
        "family": "BEBIDAS",
        "subfamily": "CACHACAS",
        "currency": "BRL",
        "source_value": Decimal("89.45"),
        "exchange_rate": Decimal("1"),
        "iof_percent": Decimal("0"),
        "cost_source": "Media online: Galeria da Cachaca R$82,50; Cachacaria Nacional R$90,29; Glouglou R$95,00; Templo da Cachaca R$90,00.",
        "data_source": "Rotulo da foto: Cachaça Matriarca Jaqueira, 750 ml, 42% vol.",
        "aliases": ["CACHACA MATRIARCA JAQUEIRA", "MATRIARCA JAQUEIRA"],
    },
    {
        "photo": "WhatsApp Image 2026-09-06 at 11.42.19(1).jpeg",
        "product_name": "CACHACA COLOMBINA TRADICIONAL 1920 JATOBA",
        "volume_ml": Decimal("700"),
        "alcohol": "41.5",
        "family": "BEBIDAS",
        "subfamily": "CACHACAS",
        "currency": "BRL",
        "source_value": Decimal("63.00"),
        "exchange_rate": Decimal("1"),
        "iof_percent": Decimal("0"),
        "cost_source": "Galeria da Cachaca: Cachaça Colombina Tradicional 700 ml a R$63,00.",
        "data_source": "Rotulo da foto e site oficial Colombina: tradicional desde 1920, 700 ml, 41,5% vol., jatoba.",
        "aliases": ["CACHACA COLOMBINA TRADICIONAL", "COLOMBINA 1920", "COLOMBINA JATOBA"],
    },
    {
        "photo": "WhatsApp Image 2026-09-06 at 11.42.19.jpeg",
        "product_name": "CACHACA YVY DESTILARIA ORIGINAL DO BRASIL",
        "volume_ml": Decimal("750"),
        "alcohol": "40",
        "family": "BEBIDAS",
        "subfamily": "CACHACAS",
        "currency": "BRL",
        "source_value": Decimal("65.00"),
        "exchange_rate": Decimal("1"),
        "iof_percent": Decimal("0"),
        "cost_source": "Valor informado pelo usuario em 2026-09-05 como Yvy Destilaria; Mercado Livre/Google Shopping indicou oferta proxima a R$67,00.",
        "data_source": "Rotulo da foto: YVY, Cachaça Destilaria Original do Brasil, 750 ml, 40% vol.; Cosmos confirma GTIN de Cachaça Yvy 750 ml.",
        "aliases": ["CACHACA YVY", "YVY DESTILARIA ORIGINAL DO BRASIL"],
        "barcode": "7898966809354",
    },
]


def normalize(value: Any) -> str:
    text = "".join(
        char for char in unicodedata.normalize("NFD", str(value or "")) if unicodedata.category(char) != "Mn"
    )
    text = re.sub(r"[^A-Za-z0-9]+", " ", text.upper()).strip()
    return re.sub(r"\s+", " ", text)


def decimal_from(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace("R$", "").replace("%", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text:
        return None
    try:
        return Decimal(text)
    except Exception:
        return None


def money(value: Decimal | None) -> str:
    if value is None:
        return ""
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)).replace(".", ",")


def number_text(value: Decimal | None) -> str:
    if value is None:
        return ""
    normalized = value.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP).normalize()
    return format(normalized, "f")


def api_json(path: str, method: str = "GET", payload: Any | None = None, params: dict[str, Any] | None = None) -> Any:
    url = f"{API_BASE}{path}"
    if params:
        url += "?" + urllib.parse.urlencode(params)
    body = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        url,
        data=body,
        method=method,
        headers={"Content-Type": "application/json"} if body is not None else {},
    )
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            raw = response.read().decode("utf-8")
            return json.loads(raw) if raw else {}
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"{method} {url} -> HTTP {error.code}: {detail}") from error


def read_price_sheet_rows() -> list[dict[str, Any]]:
    workbook = load_workbook(PRICE_WORKBOOK, data_only=True)
    rows: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
        for row_index in range(2, sheet.max_row + 1):
            row = {headers[column - 1]: sheet.cell(row_index, column).value for column in range(1, sheet.max_column + 1)}
            row["_sheet"] = sheet.title
            row["_row"] = row_index
            rows.append(row)
    return rows


def find_price_sheet_match(item: dict[str, Any], price_rows: list[dict[str, Any]]) -> dict[str, Any] | None:
    aliases = [normalize(item["product_name"]), *(normalize(alias) for alias in item.get("aliases", []))]
    for row in price_rows:
        item_name = normalize(row.get("Item"))
        if not item_name:
            continue
        for alias in aliases:
            if alias and (alias in item_name or item_name in alias):
                return row
    return None


def active_package(product: dict[str, Any]) -> dict[str, Any] | None:
    packages = product.get("packages") or []
    active = [package for package in packages if package.get("isActive") is not False]
    return active[0] if active else (packages[0] if packages else None)


def package_ml(product: dict[str, Any]) -> Decimal | None:
    package = active_package(product)
    if not package:
        return None
    quantity = decimal_from(package.get("packageQuantity"))
    if quantity is None or quantity <= 0:
        return None
    unit = normalize(package.get("packageUnit") or product.get("controlUnit"))
    if unit in {"L", "LITER", "LITRO", "LITROS"}:
        return quantity * Decimal("1000")
    return quantity


def package_cost(product: dict[str, Any]) -> Decimal | None:
    package = active_package(product)
    if not package:
        return None
    return decimal_from(package.get("purchasePrice"))


def round_up_to_5(value: Decimal) -> Decimal:
    return Decimal(str(int(math.ceil(float(value) / 5.0) * 5)))


def final_cost(item: dict[str, Any]) -> Decimal:
    converted = item["source_value"] * item["exchange_rate"]
    iof = converted * (item["iof_percent"] / Decimal("100"))
    return (converted + iof).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def suggested_sale_price(cost: Decimal, volume_ml: Decimal, product_name: str) -> Decimal:
    dose_cost = (cost / volume_ml) * DOSE_VOLUME_ML
    minimum_by_cmv = dose_cost / (MAX_CMV / Decimal("100"))
    floor = Decimal("40")
    return round_up_to_5(max(minimum_by_cmv, floor))


def find_existing_product(item: dict[str, Any], products: list[dict[str, Any]]) -> dict[str, Any] | None:
    keys = {normalize(item["product_name"]), *(normalize(alias) for alias in item.get("aliases", []))}
    for product in products:
        if product.get("technicalSheetId"):
            continue
        product_key = normalize(product.get("name"))
        if product_key in keys:
            return product
    return None


def find_existing_dose_sheet(product_id: str, sale_name: str, sheets: list[dict[str, Any]]) -> dict[str, Any] | None:
    sale_key = normalize(sale_name)
    for sheet in sorted(sheets, key=lambda record: record.get("id") or 0):
        if sheet.get("kind") != "VENDA" or sheet.get("isActive") is not True:
            continue
        if normalize(sheet.get("name")) == sale_key:
            return sheet
        for ingredient in sheet.get("ingredients") or []:
            if ingredient.get("isActive") is True and ingredient.get("productId") == product_id:
                quantity = decimal_from(ingredient.get("quantity"))
                if quantity == DOSE_VOLUME_ML:
                    return sheet
    return None


def reference_codes(item: dict[str, Any]) -> list[dict[str, Any]]:
    barcode = str(item.get("barcode") or "").strip()
    if not barcode:
        return []
    return [
        {
            "id": int(time.time() * 1000),
            "code": barcode,
            "source": "rotulo/fonte pesquisada",
            "type": "EAN" if barcode.isdigit() and len(barcode) in {8, 12, 13, 14} else "OUTRO",
            "isActive": True,
        }
    ]


def build_product_payload(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": "",
        "companyId": COMPANY_ID,
        "ownerCompanyId": COMPANY_ID,
        "companyProductId": "",
        "name": item["product_name"],
        "controlUnit": "MILLILITER",
        "family": item["family"],
        "subfamily": item["subfamily"],
        "sectors": ["BAR"],
        "alcoholPercentage": item["alcohol"],
        "densitySampleVolume": "",
        "densitySampleWeight": "",
        "ignoreStock": False,
        "excludeFromExecutionYield": False,
        "isActive": True,
        "technicalSheetId": None,
        "packages": [
            {
                "id": 1,
                "companyPackageId": "",
                "referenceCodes": reference_codes(item),
                "internalCode": "EMB-001",
                "barcode": str(item.get("barcode") or "") if str(item.get("barcode") or "").isdigit() else "",
                "packageQuantity": number_text(item["volume_ml"]),
                "packageUnit": "MILLILITER",
                "grossWeightGrams": "",
                "packagingWeightGrams": "",
                "purchasePrice": number_text(final_cost(item)),
                "openingQuantity": "",
                "isActive": True,
            }
        ],
    }


def build_sheet_payload(product: dict[str, Any], sale_name: str, sale_price: Decimal, cmv: Decimal) -> dict[str, Any]:
    ingredient_id = int(time.time() * 1000)
    return {
        "id": 0,
        "companyId": COMPANY_ID,
        "ownerCompanyId": COMPANY_ID,
        "sharedCompanyIds": [],
        "kind": "VENDA",
        "productId": "",
        "companyProductId": "",
        "companyProductIdsByCompanyId": {},
        "name": sale_name,
        "family": "DOSES 60ML",
        "subfamily": "DESTILADOS",
        "sectors": product.get("sectors") or ["BAR"],
        "outputQuantity": "1",
        "outputUnit": "UNIT",
        "densitySampleVolume": "",
        "densitySampleWeight": "",
        "yieldDifferenceDestination": "",
        "yieldDifferenceByproductName": "",
        "yieldDifferenceByproductTechnicalSheetId": None,
        "targetPh": "",
        "targetBrix": "",
        "portionSize": "1",
        "colorTagOne": "",
        "colorTagTwo": "",
        "desiredCmvPercentage": number_text(cmv),
        "dilutionRatePercentage": "",
        "imageDataUrl": "",
        "finalSalePrice": number_text(sale_price),
        "flavorProfileRatings": [],
        "flavorSweet": "0",
        "flavorSour": "0",
        "flavorBitter": "0",
        "flavorSalty": "0",
        "flavorUmami": "0",
        "storytelling": "",
        "salesArguments": "",
        "harmonization": "",
        "preparationMode": "",
        "preparationLeadTimeDays": "",
        "shelfLifeRoom": "",
        "shelfLifeRefrigerated": "",
        "shelfLifeFrozen": "",
        "productionCenters": [],
        "supplyRoutes": [],
        "ingredients": [
            {
                "id": ingredient_id,
                "productId": product["id"],
                "productLabel": product["name"],
                "quantity": number_text(DOSE_VOLUME_ML),
                "yieldQuantity": number_text(DOSE_VOLUME_ML),
                "operationalUnit": "",
                "manipulatedQuantity": "",
                "operationalQuantity": "",
                "isActive": True,
            }
        ],
        "garnishIngredients": [],
        "serviceItems": [],
        "isActive": True,
    }


def build_linked_product(sheet: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": sheet["productId"],
        "companyId": sheet["companyId"],
        "ownerCompanyId": sheet["ownerCompanyId"],
        "companyProductId": sheet.get("companyProductId") or "",
        "name": sheet["name"],
        "controlUnit": sheet["outputUnit"],
        "family": sheet.get("family") or "",
        "subfamily": sheet.get("subfamily") or "",
        "sectors": sheet.get("sectors") or [],
        "alcoholPercentage": "",
        "densitySampleVolume": "",
        "densitySampleWeight": "",
        "ignoreStock": False,
        "excludeFromExecutionYield": False,
        "isActive": True,
        "technicalSheetId": sheet["id"],
        "packages": [],
    }


def backup_online(timestamp: str) -> Path:
    target = BACKUP_ROOT / f"online-before-madre-20260906-cachacas-doses-{timestamp}"
    target.mkdir(parents=True, exist_ok=True)
    for endpoint, key, filename in [
        ("/products", "products", "products.json"),
        ("/technical-sheets", "technicalSheets", "technical-sheets.json"),
    ]:
        data = api_json(endpoint, params={"companyId": COMPANY_ID})[key]
        (target / filename).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    shutil.copy2(PRICE_WORKBOOK, target / PRICE_WORKBOOK.name)
    return target


def build_rows(products: list[dict[str, Any]], sheets: list[dict[str, Any]], price_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in PHOTO_ITEMS:
        price_sheet_match = find_price_sheet_match(item, price_rows)
        product = find_existing_product(item, products)
        item_cost = final_cost(item)
        sale_name = f"DS {item['product_name']} 60ML"
        sale_price = suggested_sale_price(item_cost, item["volume_ml"], item["product_name"])
        dose_cost = (item_cost / item["volume_ml"]) * DOSE_VOLUME_ML
        cmv = (dose_cost / sale_price) * Decimal("100")
        row = {
            **item,
            "price_sheet_match": price_sheet_match,
            "existing_product": product,
            "product_id": product.get("id") if product else "",
            "sale_name": sale_name,
            "final_cost": item_cost,
            "dose_cost": dose_cost,
            "sale_price": sale_price,
            "cmv": cmv,
            "status": "PENDENTE",
            "sheet_id": "",
            "sheet_product_id": "",
        }
        if product:
            existing_sheet = find_existing_dose_sheet(product["id"], sale_name, sheets)
            if existing_sheet:
                row["status"] = "JA_EXISTIA"
                row["sheet_id"] = existing_sheet.get("id")
                row["sheet_product_id"] = existing_sheet.get("productId")
        rows.append(row)
    return rows


def apply_rows(rows: list[dict[str, Any]], products: list[dict[str, Any]], sheets: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for row in rows:
        if row["status"] == "JA_EXISTIA":
            continue
        product = row["existing_product"]
        if not product:
            product = api_json("/products", "POST", build_product_payload(row))["product"]
            row["existing_product"] = product
            row["product_id"] = product["id"]
            row["status"] = "PRODUTO_CRIADO"
            products.append(product)
        pkg_cost = package_cost(product) or row["final_cost"]
        pkg_ml = package_ml(product) or row["volume_ml"]
        dose_cost = (pkg_cost / pkg_ml) * DOSE_VOLUME_ML
        sale_price = suggested_sale_price(pkg_cost, pkg_ml, product["name"])
        cmv = (dose_cost / sale_price) * Decimal("100")
        sale_name = f"DS {product['name']} 60ML"
        existing_sheet = find_existing_dose_sheet(product["id"], sale_name, sheets)
        if existing_sheet:
            row["status"] = f"{row['status']}_FICHA_JA_EXISTIA"
            row["sale_name"] = existing_sheet["name"]
            row["sheet_id"] = existing_sheet.get("id")
            row["sheet_product_id"] = existing_sheet.get("productId")
            continue
        saved_sheet = api_json("/technical-sheets", "POST", build_sheet_payload(product, sale_name, sale_price, cmv))["technicalSheet"]
        api_json("/products", "POST", build_linked_product(saved_sheet))
        sheets.append(saved_sheet)
        row["sale_name"] = saved_sheet["name"]
        row["sheet_id"] = saved_sheet.get("id")
        row["sheet_product_id"] = saved_sheet.get("productId")
        row["dose_cost"] = dose_cost
        row["sale_price"] = sale_price
        row["cmv"] = cmv
        row["status"] = f"{row['status']}_FICHA_CRIADA"
    return rows


def write_output_workbook(rows: list[dict[str, Any]], timestamp: str) -> None:
    if OUTPUT_WORKBOOK.exists():
        backup = OUTPUT_WORKBOOK.with_name(f"{OUTPUT_WORKBOOK.stem}.backup-{timestamp}{OUTPUT_WORKBOOK.suffix}")
        shutil.copy2(OUTPUT_WORKBOOK, backup)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Custos doses 60ml"
    headers = [
        "Foto",
        "Nome bebida",
        "ID produto base",
        "Produto ja existia",
        "Moeda origem",
        "Valor compra moeda origem",
        "Cotacao usada",
        "IOF %",
        "Valor convertido R$",
        "Valor IOF R$",
        "Outros impostos R$",
        "Outros valores inclusos R$",
        "Valor final custo R$",
        "Volume embalagem ml",
        "Custo dose 60ml R$",
        "Sugestao preco venda dose R$",
        "CMV final %",
        "ID ficha venda",
        "ID produto ficha venda",
        "Status",
        "Fonte custo",
        "Fonte dados",
        "Encontrado na Lista de preços Bar.xlsx",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row_index, row in enumerate(rows, start=2):
        ws.append(
            [
                row["photo"],
                row["product_name"],
                row["product_id"],
                "SIM" if row["existing_product"] and row["status"] == "JA_EXISTIA" else "NAO",
                row["currency"],
                float(row["source_value"]),
                float(row["exchange_rate"]),
                float(row["iof_percent"]),
                None,
                None,
                0,
                0,
                None,
                float(row["volume_ml"]),
                None,
                float(row["sale_price"]),
                None,
                row["sheet_id"],
                row["sheet_product_id"],
                row["status"],
                row["cost_source"],
                row["data_source"],
                (
                    f"{row['price_sheet_match']['_sheet']} linha {row['price_sheet_match']['_row']}: "
                    f"{row['price_sheet_match'].get('Item')} = {row['price_sheet_match'].get('Valor')}"
                    if row.get("price_sheet_match")
                    else "NAO"
                ),
            ]
        )
        ws.cell(row_index, 9, f"=F{row_index}*G{row_index}")
        ws.cell(row_index, 10, f"=I{row_index}*(H{row_index}/100)")
        ws.cell(row_index, 13, f"=I{row_index}+J{row_index}+K{row_index}+L{row_index}")
        ws.cell(row_index, 15, f"=M{row_index}/N{row_index}*60")
        ws.cell(row_index, 17, f"=O{row_index}/P{row_index}")
    for column in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(1, column).column_letter].width = min(max(len(str(ws.cell(1, column).value)) + 2, 14), 55)
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=17):
        for cell in row:
            if cell.column in {6, 7, 9, 10, 11, 12, 13, 15, 16}:
                cell.number_format = 'R$ #,##0.00'
            elif cell.column == 17:
                cell.number_format = '0.00%'
    source_ws = workbook.create_sheet("Fontes")
    source_ws.append(["Fonte", "Uso"])
    for cell in source_ws[1]:
        cell.font = Font(bold=True)
    sources = [
        ["https://www.cachacarianacional.com.br/", "Tiê Prata, Matriarca Jaqueira"],
        ["https://www.emporioincasa.com.br/", "Tiê Prata"],
        ["https://cardapio.datacaixa.com.br/loja=1596", "Tiê Prata"],
        ["https://www.magazineluiza.com.br/", "Tiê Prata"],
        ["https://fraucachaca.com/", "Tiê Castanheira, Tiê Prata referencia internacional"],
        ["https://www.cachacasdesalinas.com.br/cachaca-meia-lua-670-ml", "Meia Lua 670 ml"],
        ["https://www.galeriadacachaca.com.br/", "Meia Lua, Matriarca Jaqueira, Colombina Tradicional"],
        ["https://www.pedecuba.com.br/cachaca/armazenamento/balsamo/cachaca-meia-lua-670ml", "Meia Lua"],
        ["https://emporiodaraposa.com.br/produto/239-cachaca-com-jambu-sensacao-tradicional-750-ml", "Cachaça com Jambu Sensação"],
        ["https://www.bancadoramon.com.br/cachaca-com-jambu-sensacao-750ml", "Cachaça com Jambu Sensação"],
        ["https://pc-queijosevinhos.ola.click/", "Capueira Ouro referência online"],
        ["https://www.amazon.com.br/Famigerada-Liquor-Licor-Jenipapo-500ml/dp/B0GLSXSBQQ", "Famigerada Jeni dados do produto e referência online"],
        ["https://glouglou.com.br/", "Matriarca Jaqueira"],
        ["https://templodacachaca.com.br/", "Matriarca Jaqueira"],
        ["https://cachacacolombina.com.br/produto/colombina-3-anos/", "Colombina Tradicional dados do produto"],
        ["https://cosmos.bluesoft.com.br/produtos/7898966809354-cachaca-yvy-garrafa-750ml", "Yvy GTIN/dados do produto"],
        ["Conversa com usuario em 2026-09-05", "Custos diretos: Famigerada Jeni R$99, Capueira Ouro R$65, Yvy Destilaria R$65"],
    ]
    for source in sources:
        source_ws.append(source)
    source_ws.column_dimensions["A"].width = 80
    source_ws.column_dimensions["B"].width = 55
    workbook.save(OUTPUT_WORKBOOK)


def verify(created_rows: list[dict[str, Any]]) -> dict[str, Any]:
    products = api_json("/products", params={"companyId": COMPANY_ID})["products"]
    sheets = api_json("/technical-sheets", params={"companyId": COMPANY_ID})["technicalSheets"]
    product_by_id = {product.get("id"): product for product in products}
    sheet_by_id = {sheet.get("id"): sheet for sheet in sheets}
    missing_products = [row["product_id"] for row in created_rows if row["product_id"] not in product_by_id]
    missing_sheets = [row["sheet_id"] for row in created_rows if row["sheet_id"] not in sheet_by_id]
    bad_ingredients: list[dict[str, Any]] = []
    for row in created_rows:
        sheet = sheet_by_id.get(row["sheet_id"])
        if not sheet:
            continue
        if not any(
            ingredient.get("productId") == row["product_id"]
            and ingredient.get("isActive") is True
            and decimal_from(ingredient.get("quantity")) == DOSE_VOLUME_ML
            for ingredient in sheet.get("ingredients") or []
        ):
            bad_ingredients.append({"sheetId": row["sheet_id"], "productId": row["product_id"]})
    return {
        "productsAfter": len(products),
        "technicalSheetsAfter": len(sheets),
        "missingProducts": missing_products,
        "missingSheets": missing_sheets,
        "badIngredients": bad_ingredients,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    timestamp = time.strftime("%Y%m%d-%H%M%S")
    products = api_json("/products", params={"companyId": COMPANY_ID})["products"]
    sheets = api_json("/technical-sheets", params={"companyId": COMPANY_ID})["technicalSheets"]
    price_rows = read_price_sheet_rows()
    rows = build_rows(products, sheets, price_rows)

    if args.apply:
        backup_dir = backup_online(timestamp)
        rows = apply_rows(rows, products, sheets)
    else:
        backup_dir = None
        for row in rows:
            if row["status"] == "PENDENTE":
                row["status"] = "CRIARIA_PRODUTO_E_FICHA"

    write_output_workbook(rows, timestamp)
    AUDIT_ROOT.mkdir(parents=True, exist_ok=True)
    audit_path = AUDIT_ROOT / f"madre-20260906-photo-cachacas-doses-{timestamp}.json"
    audit = {
        "apply": args.apply,
        "backupDir": str(backup_dir) if backup_dir else "",
        "outputWorkbook": str(OUTPUT_WORKBOOK),
        "rows": rows,
        "verification": verify(rows) if args.apply else {},
    }
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
